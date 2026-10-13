import os
import re
from datetime import date, datetime
from pathlib import Path
from sqlalchemy.orm import Session
from sqlalchemy import case, func
from app.models.reports import Report
from app.models.feedback import Feedback
from app.models.detection import DetectionLog
from app.models.announcement import Announcement
from app.models.audit import AuditLog
from app.models.user import User, Admin, Official, Resident
from app.paths import REPORT_UPLOAD_DIR

IMAGE_EXTS = {".jpg", ".jpeg", ".png", ".gif", ".bmp"}
MAX_CONTENT_LEN = 80


def _truncate(text, length=MAX_CONTENT_LEN):
    if not text:
        return ""
    s = str(text)
    return s[:length] + "..." if len(s) > length else s


def _resolve_user_name(db: Session, user_id):
    if not user_id:
        return "—"
    user = db.query(User).filter(User.user_id == user_id).first()
    if not user:
        return str(user_id)
    for profile in (user.admin_profile, user.official_profile, user.resident_profile):
        if profile:
            parts = [profile.first_name, profile.last_name]
            return " ".join(p for p in parts if p)
    return user.username or str(user_id)


def _get_attachment_info(record):
    path = getattr(record, "attachment_path", None)
    if not path:
        return None, None
    full = path.lstrip("/")
    if not os.path.exists(full):
        return None, None
    ext = Path(full).suffix.lower()
    name = Path(full).name
    return full, ext


def _get_columns(report_type):
    if report_type == "announcement":
        return ["#", "Title", "Content", "Date Posted", "Published", "Author"]
    elif report_type == "detection":
        return ["#", "Camera", "Event Type", "Confidence", "Timestamp"]
    elif report_type == "feedback":
        return ["#", "Subject", "Content", "Resolved", "Date"]
    elif report_type == "activity":
        return ["#", "Action", "Target", "Description", "User", "Timestamp"]
    return []


def _get_rows(db, report_type, records):
    rows = []
    for i, rec in enumerate(records, 1):
        if report_type == "announcement":
            author_name = _resolve_user_name(db, rec.created_by)
            rows.append([
                str(i),
                _truncate(rec.title, 30),
                _truncate(rec.content, 60),
                str(rec.date_posted.strftime("%Y-%m-%d %H:%M")) if rec.date_posted else "—",
                "Yes" if rec.is_published else "No",
                author_name,
            ])
        elif report_type == "detection":
            rows.append([
                str(i),
                rec.camera_name or str(rec.camera_id) or "—",
                rec.event_type or "—",
                f"{rec.confidence_score:.1f}%" if rec.confidence_score is not None else "—",
                str(rec.timestamp.strftime("%Y-%m-%d %H:%M")) if rec.timestamp else "—",
            ])
        elif report_type == "feedback":
            author_name = _resolve_user_name(db, rec.created_by)
            rows.append([
                str(i),
                _truncate(rec.subject, 30),
                _truncate(rec.content, 60),
                "Yes" if rec.is_resolved else "No",
                str(rec.timestamp.strftime("%Y-%m-%d %H:%M")) if rec.timestamp else "—",
            ])
        elif report_type == "activity":
            user_name = _resolve_user_name(db, rec.user_id)
            rows.append([
                str(i),
                _truncate(rec.action_type, 30),
                rec.target_table or "—",
                _truncate(rec.description, 40),
                user_name,
                str(rec.timestamp.strftime("%Y-%m-%d %H:%M")) if rec.timestamp else "—",
            ])
    return rows


def _generate_pdf(filepath, title, report_type, start_date, end_date, columns, rows, total, records, db):
    from reportlab.lib.pagesizes import letter
    from reportlab.lib import colors
    from reportlab.lib.units import inch
    from reportlab.platypus import (
        SimpleDocTemplate, Table, TableStyle, Paragraph, Spacer, Image as RLImage,
        PageBreak,
    )
    from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
    from reportlab.lib.enums import TA_CENTER, TA_LEFT

    styles = getSampleStyleSheet()
    title_style = ParagraphStyle("ReportTitle", parent=styles["Title"], fontSize=20, spaceAfter=6)
    subtitle_style = ParagraphStyle("ReportSub", parent=styles["Normal"], fontSize=11, textColor=colors.grey, alignment=TA_CENTER, spaceAfter=4)
    cell_style = ParagraphStyle("Cell", parent=styles["Normal"], fontSize=8, leading=10, wordWrap="CJK")
    cell_bold = ParagraphStyle("CellBold", parent=cell_style, fontName="Helvetica-Bold")

    def add_page_number(canvas, doc):
        canvas.saveState()
        canvas.setFont("Helvetica", 9)
        canvas.drawCentredString(letter[0] / 2, 30, f"Page {canvas.getPageNumber()}")
        canvas.restoreState()

    doc = SimpleDocTemplate(str(filepath), pagesize=letter,
                            leftMargin=50, rightMargin=50, topMargin=60, bottomMargin=60)
    elements = []

    elements.append(Paragraph(title, title_style))
    elements.append(Spacer(1, 12))
    elements.append(Paragraph(f"Report Type: {report_type.upper()}", subtitle_style))
    elements.append(Paragraph(f"Date Range: {start_date or 'All'} to {end_date or 'All'}", subtitle_style))
    elements.append(Paragraph(f"Total Records: {total}", subtitle_style))
    elements.append(Spacer(1, 30))

    if not rows:
        elements.append(Paragraph("No records found for the selected criteria.", styles["Normal"]))
        doc.build(elements, onFirstPage=add_page_number, onLaterPages=add_page_number)
        return

    header = [Paragraph(f"<b>{c}</b>", cell_bold) for c in columns]
    table_data = [header]
    for row in rows:
        table_data.append([Paragraph(str(cell), cell_style) for cell in row])

    page_width = letter[0] - 100
    num_cols = len(columns)
    col_width = page_width / num_cols
    col_widths = [col_width] * num_cols

    table = Table(table_data, colWidths=col_widths, repeatRows=1)
    style_cmds = [
        ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#1a237e")),
        ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
        ("FONTNAME", (0, 0), (-1, 0), "Helvetica-Bold"),
        ("FONTSIZE", (0, 0), (-1, 0), 9),
        ("BOTTOMPADDING", (0, 0), (-1, 0), 8),
        ("TOPPADDING", (0, 0), (-1, 0), 8),
        ("ALIGN", (0, 0), (-1, -1), "LEFT"),
        ("VALIGN", (0, 0), (-1, -1), "TOP"),
        ("GRID", (0, 0), (-1, -1), 0.5, colors.grey),
        ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.white, colors.HexColor("#f5f5f5")]),
        ("TOPPADDING", (0, 1), (-1, -1), 4),
        ("BOTTOMPADDING", (0, 1), (-1, -1), 4),
        ("LEFTPADDING", (0, 0), (-1, -1), 6),
        ("RIGHTPADDING", (0, 0), (-1, -1), 6),
    ]
    table.setStyle(TableStyle(style_cmds))
    elements.append(table)

    if report_type in ("announcement", "feedback"):
        elements.append(Spacer(1, 20))
        elements.append(Paragraph("Attachments", styles["Heading2"]))
        elements.append(Spacer(1, 8))
        has_attachments = False
        for i, rec in enumerate(records, 1):
            full, ext = _get_attachment_info(rec) if report_type == "announcement" else (None, None)
            if report_type == "feedback":
                full, ext = _get_attachment_info(rec)
            if full:
                has_attachments = True
                label = f"Record #{i}"
                elements.append(Paragraph(f"<b>{label}</b>", styles["Normal"]))
                elements.append(Spacer(1, 4))
                if ext in IMAGE_EXTS:
                    try:
                        img = RLImage(full, width=4 * inch, height=3 * inch)
                        elements.append(img)
                    except Exception:
                        elements.append(Paragraph(f"Attachment: {Path(full).name}", styles["Normal"]))
                else:
                    elements.append(Paragraph(f"Attachment: {Path(full).name}", styles["Normal"]))
                elements.append(Spacer(1, 12))
        if not has_attachments:
            elements.append(Paragraph("No attachments found.", styles["Normal"]))

    doc.build(elements, onFirstPage=add_page_number, onLaterPages=add_page_number)


def _generate_docx(filepath, title, report_type, start_date, end_date, columns, rows, total, records, db):
    from docx import Document
    from docx.shared import Inches, Pt, RGBColor
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.enum.table import WD_TABLE_ALIGNMENT
    from docx.oxml.ns import qn, nsdecls
    from docx.oxml import parse_xml

    doc = Document()

    heading = doc.add_heading(title, level=0)
    heading.alignment = WD_ALIGN_PARAGRAPH.CENTER

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run(f"Report Type: {report_type.upper()}")
    run.font.size = Pt(11)
    run.font.color.rgb = RGBColor(100, 100, 100)

    p2 = doc.add_paragraph()
    p2.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run2 = p2.add_run(f"Date Range: {start_date or 'All'} to {end_date or 'All'} | Total Records: {total}")
    run2.font.size = Pt(10)
    run2.font.color.rgb = RGBColor(100, 100, 100)

    doc.add_paragraph("")

    if not rows:
        doc.add_paragraph("No records found for the selected criteria.")
        doc.save(str(filepath))
        return

    table = doc.add_table(rows=1, cols=len(columns))
    table.style = "Table Grid"
    table.alignment = WD_TABLE_ALIGNMENT.CENTER

    hdr = table.rows[0]
    for idx, col_name in enumerate(columns):
        cell = hdr.cells[idx]
        cell.text = ""
        run = cell.paragraphs[0].add_run(col_name)
        run.bold = True
        run.font.size = Pt(9)
        run.font.color.rgb = RGBColor(255, 255, 255)
        shading = parse_xml(f'<w:shd {nsdecls("w")} w:fill="1a237e"/>')
        cell._tc.get_or_add_tcPr().append(shading)

    for row_data in rows:
        row_cells = table.add_row().cells
        for idx, val in enumerate(row_data):
            cell = row_cells[idx]
            cell.text = ""
            run = cell.paragraphs[0].add_run(str(val))
            run.font.size = Pt(8)

    if report_type in ("announcement", "feedback"):
        doc.add_page_break()
        doc.add_heading("Attachments", level=1)
        has_attachments = False
        for i, rec in enumerate(records, 1):
            full, ext = _get_attachment_info(rec)
            if full:
                has_attachments = True
                doc.add_paragraph(f"Record #{i}", style="Heading3")
                if ext in IMAGE_EXTS:
                    try:
                        doc.add_picture(full, width=Inches(5))
                    except Exception:
                        doc.add_paragraph(f"Attachment: {Path(full).name}")
                else:
                    doc.add_paragraph(f"Attachment: {Path(full).name}")
        if not has_attachments:
            doc.add_paragraph("No attachments found.")

    doc.save(str(filepath))


def generate_report(
    db: Session,
    title: str,
    report_type: str,
    start_date: date | None = None,
    end_date: date | None = None,
    file_format: str = "pdf",
    created_by: int | None = None,
) -> Report:
    records = []
    if report_type == "feedback":
        q = db.query(Feedback)
        if start_date:
            q = q.filter(Feedback.timestamp >= start_date)
        if end_date:
            q = q.filter(Feedback.timestamp <= end_date)
        records = q.order_by(Feedback.timestamp.desc()).all()
    elif report_type == "detection":
        q = db.query(DetectionLog)
        if start_date:
            q = q.filter(DetectionLog.timestamp >= start_date)
        if end_date:
            q = q.filter(DetectionLog.timestamp <= end_date)
        records = q.order_by(DetectionLog.timestamp.desc()).all()
    elif report_type == "announcement":
        q = db.query(Announcement)
        if start_date:
            q = q.filter(Announcement.date_posted >= start_date)
        if end_date:
            q = q.filter(Announcement.date_posted <= end_date)
        records = q.order_by(Announcement.date_posted.desc()).all()
    elif report_type == "activity":
        q = db.query(AuditLog)
        if start_date:
            q = q.filter(AuditLog.timestamp >= start_date)
        if end_date:
            q = q.filter(AuditLog.timestamp <= end_date)
        records = q.order_by(AuditLog.timestamp.desc()).all()

    safe_title = re.sub(r'[^\w\s\-]', '', title).strip().replace(' ', '_') or report_type
    filename = f"{safe_title}_{datetime.now().strftime('%Y%m%d_%H%M%S')}.{file_format}"
    filepath = REPORT_UPLOAD_DIR / filename
    filepath.parent.mkdir(parents=True, exist_ok=True)

    columns = _get_columns(report_type)
    rows = _get_rows(db, report_type, records)
    total = len(records)

    if file_format == "docx":
        _generate_docx(filepath, title, report_type, start_date, end_date, columns, rows, total, records, db)
    else:
        _generate_pdf(filepath, title, report_type, start_date, end_date, columns, rows, total, records, db)

    report = Report(
        title=title,
        file_format=file_format,
        report_type=report_type,
        start_date=start_date,
        end_date=end_date,
        file_path=str(filepath),
        created_by=created_by,
    )
    db.add(report)
    db.commit()
    db.refresh(report)
    return report
